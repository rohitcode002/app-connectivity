"""
cmets_handler/formatter.py — Final CMETS value formatter
========================================================
Creates a cleaned CMETS workbook after extraction/consolidation and before
downstream mapping consumes the CMETS data.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from pipeline.excel_utils import (
    _apply_data_style,
    _apply_header_style,
    _autosize_columns,
    _get_openpyxl,
)
from pipeline.cmets_handler.normalization import INDIA_STATES_UTS, clean, norm_substation
from pipeline.cmets_handler.normalization import extract_date, gna_yes_no
from pipeline.cmets_handler.normalization import norm_type, parse_type_capacity


CAPACITY_COLUMNS = [
    "Installed/Break-up Capacity (MW) Solar",
    "Installed/Break-up Capacity (MW) Wind",
    "Installed/Break-up Capacity (MW) Hybrid",
    "Installed/Break-up Capacity (MW) Hydro",
]

REGION_COLUMN = "Region"

FINAL_FORMATTED_COLUMNS = [
    "Region",
    "State",
    "Substation",
    "Coordinates",
    "Project Location",
    "Name of Developers",
    "GNA/ST II Application ID",
    "LTA Application ID",
    "Application ID under Enhancement 5.2 or revision",
    "CMETS GNA Approved",
    "CMETS LTA Approved",
    "CMETS GNA Meeting Date",
    "CMETS LTA Meeting Date",
    "Type",
    "Application Quantum (MW)(ST II)",
    "Granted Quantum GNA/LTA(MW)",
    "Installed/Break-up Capacity (MW) Solar",
    "Installed/Break-up Capacity (MW) Wind",
    "Installed/Break-up Capacity (MW) Hybrid",
    "Installed/Break-up Capacity (MW) Hydro",
    "Battery MWh",
    "Battery Injection (MW)",
    "Battery Drawl (MW)",
    "PSP MWh",
    "PSP Injection (MW)",
    "PSP Drawl (MW)",
    "Commissioned TGNA",
    "Commissioned GNA",
    "Application/Submission Date",
    "Mode(Criteria for applying)",
    "Applied Start of Connectivity sought by developer date"
    "( start date of connectivity as per the application)",
    "GNA Operationalization Date",
    "GNA Operationalization (Yes/No)",
    "Date from which additional capacity is to be added",
    "Nature of Applicant",
    "Status of application(Withdrawn / granted. Revoked.)",
    "Voltage level",
    "Bay No",
]

_REGION_PATTERNS = [
    (r"\bnorth[\s_-]*eastern[\s_-]*region\b", "NER"),
    (r"\bnorth[\s_-]*east(?:ern)?\b", "NER"),
    (r"\bnorthern[\s_-]*region\b", "NR"),
    (r"\bsouthern[\s_-]*region\b", "SR"),
    (r"\bwestern[\s_-]*region\b", "WR"),
    (r"\beastern[\s_-]*region\b", "ER"),
    (r"\bCMETS[\s_-]*NER\b", "NER"),
    (r"\bCMETS[\s_-]*NR\b", "NR"),
    (r"\bCMETS[\s_-]*SR\b", "SR"),
    (r"\bCMETS[\s_-]*WR\b", "WR"),
    (r"\bCMETS[\s_-]*ER\b", "ER"),
    (r"\bNER\b", "NER"),
    (r"\bNR\b", "NR"),
    (r"\bSR\b", "SR"),
    (r"\bWR\b", "WR"),
    (r"\bER\b", "ER"),
]

ID_COLUMNS = [
    "GNA/ST II Application ID",
    "LTA Application ID",
    "Application ID under Enhancement 5.2 or revision",
]

DATE_COLUMNS = [
    "CMETS GNA Meeting Date",
    "CMETS LTA Meeting Date",
    "Application/Submission Date",
    "Applied Start of Connectivity sought by developer date"
    "( start date of connectivity as per the application)",
    "GNA Operationalization Date",
    "Date from which additional capacity is to be added",
]


def _text(value: Any) -> str:
    value = clean(value)
    return str(value) if value is not None else ""


def _numbers_only(value: Any, *, min_digits: int = 1) -> str | None:
    pattern = r"\d+" if min_digits <= 1 else rf"\b\d{{{min_digits},}}\b"
    nums = re.findall(pattern, _text(value))
    return ", ".join(nums) if nums else None


def _format_date(value: Any) -> str | None:
    return extract_date(_text(value))


def _format_state(value: Any) -> str | None:
    text = _text(value).lower()
    if not text:
        return None
    for state in sorted(INDIA_STATES_UTS, key=len, reverse=True):
        if state in text:
            return state.title()
    return text.split(",")[-1].strip(" .").title() if "," in text else text.title()


def _format_region(*values: Any) -> str | None:
    text = " ".join(_text(value) for value in values if _text(value))
    if not text:
        return None
    for pattern, region in _REGION_PATTERNS:
        if re.search(pattern, text, flags=re.IGNORECASE):
            return region
    return None


def _format_substation(value: Any) -> str | None:
    return norm_substation(_text(value))


def _format_status(value: Any) -> str | None:
    text = _text(value).lower()
    if not text:
        return "Applied"
    if "withdraw" in text:
        return "Withdrawn"
    if any(word in text for word in ("revoke", "cancel", "reject")):
        return "Revoked"
    if "grant" in text or "approved" in text:
        return "granted"
    return "Applied"


_HYBRID_VALUE_RE = re.compile(
    r"\b(?:hybrid|solar\s*\+\s*wind)\s*\(\s*([\d,.]+)",
    re.IGNORECASE,
)


def _format_number(value: float) -> int | float:
    return int(value) if float(value).is_integer() else round(value, 4)


def _format_numeric_cell(value: Any) -> int | float | None:
    match = re.search(r"\d+(?:,\d{3})*(?:\.\d+)?", _text(value))
    if not match:
        return None
    return _format_number(float(match.group(0).replace(",", "")))


def _parse_type(value: Any) -> tuple[str | None, dict[str, int | float]]:
    text = _text(value)
    if not text:
        return None, {}

    component_caps = parse_type_capacity(text)
    capacities = {"solar": 0.0, "wind": 0.0, "hybrid": 0.0, "hydro": 0.0}
    capacities["solar"] = component_caps.get("Solar", 0.0)
    capacities["wind"] = component_caps.get("Wind", 0.0)
    capacities["hydro"] = component_caps.get("Hydro", 0.0)

    for match in _HYBRID_VALUE_RE.finditer(text):
        capacities["hybrid"] += float(match.group(1).replace(",", ""))

    formatted_type = norm_type(text)
    formatted_caps = {
        "Installed/Break-up Capacity (MW) Solar": _format_number(capacities["solar"]),
        "Installed/Break-up Capacity (MW) Wind": _format_number(capacities["wind"]),
        "Installed/Break-up Capacity (MW) Hybrid": _format_number(capacities["hybrid"]),
        "Installed/Break-up Capacity (MW) Hydro": _format_number(capacities["hydro"]),
    }
    formatted_caps = {k: v for k, v in formatted_caps.items() if v}
    return formatted_type, formatted_caps


def _headers(ws) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in ws[1] if cell.value}


def _ensure_columns(ws, columns: list[str]) -> dict[str, int]:
    headers = _headers(ws)
    for col in columns:
        if col not in headers:
            ws.cell(row=1, column=ws.max_column + 1, value=col)
            headers[col] = ws.max_column
    return headers


def _set(ws, row_idx: int, headers: dict[str, int], col: str, value: Any) -> None:
    if col in headers:
        ws.cell(row=row_idx, column=headers[col], value=value)


def _get(ws, row_idx: int, headers: dict[str, int], col: str) -> Any:
    if col not in headers:
        return None
    return ws.cell(row=row_idx, column=headers[col]).value


def _reorder_columns(ws, ordered_columns: list[str]) -> None:
    """Rewrite the worksheet with exactly the requested user-facing columns."""
    headers = _headers(ws)
    rows: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        rows.append({
            col: ws.cell(row=row_idx, column=col_idx).value
            for col, col_idx in headers.items()
        })

    if ws.max_column:
        ws.delete_cols(1, ws.max_column)

    for col_idx, col_name in enumerate(ordered_columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, col_name in enumerate(ordered_columns, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col_name))


def format_cmets_excel(source_path: str | Path, output_path: str | Path | None = None) -> Path:
    """Create a final formatted CMETS workbook and return its path."""
    src = Path(source_path).resolve()
    out = Path(output_path).resolve() if output_path else src.with_name(f"{src.stem}_final_formatted.xlsx")

    wb = load_workbook(src)
    ws = wb["Extracted Data"] if "Extracted Data" in wb.sheetnames else wb.active
    headers = _ensure_columns(ws, FINAL_FORMATTED_COLUMNS)

    for row_idx in range(2, ws.max_row + 1):
        region = _format_region(
            _get(ws, row_idx, headers, REGION_COLUMN),
            _get(ws, row_idx, headers, "PDF"),
        )
        _set(ws, row_idx, headers, REGION_COLUMN, region)

        state = _format_state(_get(ws, row_idx, headers, "State")) or _format_state(
            _get(ws, row_idx, headers, "Project Location")
        )
        _set(ws, row_idx, headers, "State", state)
        _set(ws, row_idx, headers, "Substation", _format_substation(_get(ws, row_idx, headers, "Substation")))

        for col in ID_COLUMNS:
            _set(ws, row_idx, headers, col, _numbers_only(_get(ws, row_idx, headers, col), min_digits=6))

        _set(ws, row_idx, headers, "CMETS GNA Approved", _numbers_only(_get(ws, row_idx, headers, "CMETS GNA Approved")))
        _set(ws, row_idx, headers, "CMETS LTA Approved", _numbers_only(_get(ws, row_idx, headers, "CMETS LTA Approved")))

        for col in DATE_COLUMNS:
            _set(ws, row_idx, headers, col, _format_date(_get(ws, row_idx, headers, col)))

        _set(
            ws,
            row_idx,
            headers,
            "Application Quantum (MW)(ST II)",
            _format_numeric_cell(_get(ws, row_idx, headers, "Application Quantum (MW)(ST II)")),
        )

        _set(
            ws,
            row_idx,
            headers,
            "GNA Operationalization (Yes/No)",
            gna_yes_no(_get(ws, row_idx, headers, "GNA Operationalization Date")),
        )
        _set(
            ws,
            row_idx,
            headers,
            "Status of application(Withdrawn / granted. Revoked.)",
            _format_status(_get(ws, row_idx, headers, "Status of application(Withdrawn / granted. Revoked.)")),
        )

        # ── Granted Quantum: copy Application Quantum when status is granted ─
        status_val = _text(_get(ws, row_idx, headers, "Status of application(Withdrawn / granted. Revoked.)")).lower()
        if status_val == "granted":
            _set(ws, row_idx, headers, "Granted Quantum GNA/LTA(MW)",
                 _format_numeric_cell(_get(ws, row_idx, headers, "Application Quantum (MW)(ST II)")))
        else:
            _set(ws, row_idx, headers, "Granted Quantum GNA/LTA(MW)", None)

        formatted_type, capacities = _parse_type(_get(ws, row_idx, headers, "Type"))
        _set(ws, row_idx, headers, "Type", formatted_type)
        for col in CAPACITY_COLUMNS:
            _set(ws, row_idx, headers, col, capacities.get(col))

    _reorder_columns(ws, FINAL_FORMATTED_COLUMNS)

    opx = _get_openpyxl()
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _apply_header_style(ws, opx)
    _apply_data_style(ws, opx)
    _autosize_columns(ws)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
